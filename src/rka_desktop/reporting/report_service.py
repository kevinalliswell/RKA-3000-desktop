from __future__ import annotations

from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from rka_desktop.analysis.reduction_kinetics import ReductionResult


class ReportService:
    """Generate technical experiment records and the existing kinetics workbook."""

    def render_record_html(
        self,
        experiment: dict[str, Any],
        measurements: list[dict[str, Any]],
        alarms: list[dict[str, Any]],
    ) -> str:
        """Render a self-contained, escaped record. It makes no method-compliance claim."""
        def cell(value: Any) -> str:
            return escape("—" if value is None or value == "" else str(value), quote=True)

        def number(value: Any, decimals: int = 3) -> str:
            try:
                return f"{float(value):.{decimals}f}"
            except (TypeError, ValueError):
                return "—"

        generated_at = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
        status = str(experiment.get("status") or "UNKNOWN")
        warnings = ["本文件为技术实验记录，尚未经方法符合性验证或审核签署。"]
        if status != "FINISHED":
            warnings.append(f"实验状态为 {status}，记录可能不完整。")
        if not measurements:
            warnings.append("无采样数据，不能据此作出实验结果判断。")
        if alarms:
            warnings.append(f"存在 {len(alarms)} 条报警记录，须结合现场处置复核。")
        warnings.append("当前数据库未保存启动时的完整配方快照，报告仅列出配方关联名称。")

        meta = (
            ("实验编号", experiment.get("id")), ("实验名称", experiment.get("name")),
            ("实验状态", status), ("关联配方", experiment.get("recipe_name")),
            ("开始时间", experiment.get("started_at")), ("结束时间", experiment.get("finished_at")),
            ("初始质量 m₀ / g", number(experiment.get("initial_mass_g"))),
            ("全铁 TFe / %", number(experiment.get("tfe_pct"), 2)),
            ("FeO / %", number(experiment.get("feo_pct"), 2)),
            ("采样数", len(measurements)), ("报警数", len(alarms)),
        )
        meta_rows = "".join(f"<tr><th>{cell(label)}</th><td>{cell(value)}</td></tr>" for label, value in meta)
        warning_rows = "".join(f"<li>{cell(message)}</li>" for message in warnings)
        sample_rows = "".join(
            "<tr>"
            + f"<td>{cell(row.get('sampled_at'))}</td>"
            + "".join(f"<td class='num'>{number(row.get(key), digits)}</td>" for key, digits in (
                ("temperature_c", 2), ("weight_g", 3), ("pressure_mpa", 4),
                ("n2_lpm", 3), ("h2_lpm", 3), ("co_lpm", 3),
            ))
            + "</tr>"
            for row in measurements
        ) or "<tr><td colspan='7'>无采样数据</td></tr>"
        alarm_rows = "".join(
            "<tr>" + "".join(f"<td>{cell(alarm.get(key))}</td>" for key in (
                "created_at", "code", "level", "message", "acknowledged_at",
            )) + "</tr>"
            for alarm in alarms
        ) or "<tr><td colspan='5'>无报警记录</td></tr>"
        return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>RKA-3000 实验记录 #{cell(experiment.get('id'))}</title>
<style>
@page {{ size: A4; margin: 18mm; }}
body {{ font-family: 'PingFang SC','Microsoft YaHei',Arial,sans-serif; color:#18202f; margin:32px auto; max-width:1000px; font-size:10pt; line-height:1.5; }}
header {{ border-bottom:3px solid #0369a1; padding-bottom:14px; margin-bottom:20px; }}
h1 {{ margin:4px 0; font-size:22px; }} h2 {{ color:#075985; font-size:16px; margin-top:28px; }}
.eyebrow {{ color:#0369a1; font-weight:700; letter-spacing:1px; font-size:9pt; }}
.badge {{ display:inline-block; background:#fff3d6; color:#8a5a00; padding:4px 8px; border-radius:4px; font-size:9pt; }}
table {{ border-collapse:collapse; width:100%; margin:10px 0; table-layout:fixed; }}
th,td {{ border:1px solid #d9e1eb; padding:6px 8px; text-align:left; word-break:break-word; font-size:9pt; }}
th {{ background:#edf2f7; }} .meta th {{ width:26%; }} .num {{ text-align:right; font-variant-numeric:tabular-nums; }}
.notice {{ background:#fff8e6; border-left:3px solid #d18b00; padding:8px 14px; }}
.small {{ color:#617084; font-size:9pt; }}
tr {{ break-inside:avoid; }}
</style></head><body>
<header><div class="eyebrow">RKA-3000 / EXPERIMENT RECORD</div><h1>铁矿石氧化还原实验记录</h1>
<span class="badge">技术记录 · 待复核</span><p class="small">生成时间：{cell(generated_at)}</p></header>
<h2>一、实验信息</h2><table class="meta" width="100%" cellspacing="0" cellpadding="6">{meta_rows}</table>
<h2>二、数据与方法限制</h2><div class="notice"><ul>{warning_rows}</ul></div>
<h2>三、报警记录</h2><table width="100%" cellspacing="0" cellpadding="6"><thead><tr><th>时间</th><th>编码</th><th>级别</th><th>内容</th><th>确认时间</th></tr></thead><tbody>{alarm_rows}</tbody></table>
<h2>四、原始采样</h2><p class="small">按数据库保存顺序列出；完整性需结合采样间隔和设备状态复核。</p>
<table width="100%" cellspacing="0" cellpadding="6"><thead><tr><th>采样时间</th><th>温度 ℃</th><th>质量 g</th><th>压力 MPa</th><th>N₂ L/min</th><th>H₂ L/min</th><th>CO L/min</th></tr></thead><tbody>{sample_rows}</tbody></table>
<p class="small">本文件不包含正式审签；还原度分析另由分析页面计算并导出。</p>
</body></html>"""

    def generate_html_record(
        self,
        experiment: dict[str, Any],
        measurements: list[dict[str, Any]],
        alarms: list[dict[str, Any]],
        file_path: Path,
    ) -> Path:
        file_path.parent.mkdir(parents=True, exist_ok=True)
        file_path.write_text(self.render_record_html(experiment, measurements, alarms), encoding="utf-8")
        return file_path

    def generate_pdf_record(
        self,
        experiment: dict[str, Any],
        measurements: list[dict[str, Any]],
        alarms: list[dict[str, Any]],
        file_path: Path,
    ) -> Path:
        """Print the same technical record to PDF using Qt already bundled with the app."""
        from PySide6.QtGui import QTextDocument
        from PySide6.QtPrintSupport import QPrinter

        file_path.parent.mkdir(parents=True, exist_ok=True)
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(str(file_path))
        document = QTextDocument()
        document.setHtml(self.render_record_html(experiment, measurements, alarms))
        document.print_(printer)
        if not file_path.is_file() or file_path.stat().st_size == 0:
            raise OSError("PDF 报告未成功生成")
        return file_path

    def generate_excel_report(
        self,
        experiment: dict[str, Any],
        result: ReductionResult,
        file_path: Path,
    ) -> Path:
        file_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()

        summary = wb.active
        summary.title = "报告"
        summary["A1"] = "铁矿石还原动力学实验报告（失重法）"
        rows = [
            ("实验编号", experiment.get("id")),
            ("实验名称", experiment.get("name")),
            ("状态", experiment.get("status")),
            ("初始质量 m0 (g)", experiment.get("initial_mass_g")),
            ("全铁 TFe (%)", experiment.get("tfe_pct")),
            ("FeO (%)", experiment.get("feo_pct")),
            ("终态还原度 (%)", round(result.final_degree_pct, 3)),
            ("到 40% 时间 (min)", result.time_to_40_min),
            ("到 60% 时间 (min)", result.time_to_60_min),
            ("40% 还原速率 (%/min)", result.rate_at_40_pct),
        ]
        for i, (label, value) in enumerate(rows, start=3):
            summary[f"A{i}"] = label
            summary[f"B{i}"] = value

        data = wb.create_sheet("数据")
        data.append(["时间 (min)", "还原度 R (%)", "还原速率 dR/dt (%/min)"])
        for point in result.points:
            data.append(
                [
                    round(point.t_min, 4),
                    round(point.reduction_degree_pct, 4),
                    round(point.rate_pct_per_min, 4),
                ]
            )

        n = len(result.points)
        if n >= 2:
            chart = LineChart()
            chart.title = "还原度曲线 R(t)"
            chart.x_axis.title = "时间 (min)"
            chart.y_axis.title = "还原度 (%)"
            values = Reference(data, min_col=2, min_row=1, max_row=n + 1)
            cats = Reference(data, min_col=1, min_row=2, max_row=n + 1)
            chart.add_data(values, titles_from_data=True)
            chart.set_categories(cats)
            data.add_chart(chart, "E2")

        wb.save(file_path)
        return file_path
