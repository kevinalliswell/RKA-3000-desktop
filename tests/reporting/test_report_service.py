import pytest
from openpyxl import load_workbook

from rka_desktop.analysis.reduction_kinetics import SampleParams, compute_reduction
from rka_desktop.reporting.report_service import ReportService


def _result():
    params = SampleParams(initial_mass_g=50.0, tfe_pct=60.0, feo_pct=0.0)
    series = [(0.0, 50.0), (60.0, 43.55), (120.0, 37.1)]
    return compute_reduction(series, params)


def test_generates_excel_report_with_sheets_and_chart(tmp_path):
    experiment = {
        "id": 7,
        "name": "exp-7",
        "status": "FINISHED",
        "initial_mass_g": 50.0,
        "tfe_pct": 60.0,
        "feo_pct": 0.0,
    }
    out = tmp_path / "report.xlsx"
    ReportService().generate_excel_report(experiment, _result(), out)

    assert out.exists()
    wb = load_workbook(out)
    assert set(wb.sheetnames) == {"报告", "数据"}

    report = wb["报告"]
    # indices written to the summary sheet
    labels = {report[f"A{i}"].value for i in range(3, 13)}
    assert "终态还原度 (%)" in labels
    assert "40% 还原速率 (%/min)" in labels

    data = wb["数据"]
    assert data["A1"].value == "时间 (min)"
    assert data.max_row == 4  # header + 3 points
    # a chart was embedded
    assert len(data._charts) == 1


def test_report_handles_single_point_without_chart(tmp_path):
    params = SampleParams(initial_mass_g=50.0, tfe_pct=60.0, feo_pct=0.0)
    result = compute_reduction([(0.0, 50.0)], params)
    experiment = {"id": 1, "name": "x", "status": "RUNNING"}
    out = tmp_path / "r.xlsx"
    ReportService().generate_excel_report(experiment, result, out)
    wb = load_workbook(out)
    assert len(wb["数据"]._charts) == 0  # not enough points for a chart


def test_html_record_escapes_fields_and_preserves_raw_samples(tmp_path):
    service = ReportService()
    experiment = {"id": 9, "name": "<script>alert(1)</script>", "status": "FINISHED"}
    measurements = [{
        "sampled_at": "2026-10-28T10:00:00", "temperature_c": 900,
        "weight_g": 54.2, "pressure_mpa": 0.1,
        "n2_lpm": 4, "h2_lpm": 1, "co_lpm": 0.5,
    }]
    alarms = [{"created_at": "2026-10-28T10:01:00", "code": "A1", "level": "WARN",
               "message": "<unsafe>", "acknowledged_at": None}]
    output = tmp_path / "record.html"
    service.generate_html_record(experiment, measurements, alarms, output)

    html = output.read_text(encoding="utf-8")
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<script>" not in html
    assert "&lt;unsafe&gt;" in html
    assert "54.200" in html
    assert "技术记录 · 待复核" in html


def test_pdf_record_can_be_generated(tmp_path):
    QApplication = pytest.importorskip("PySide6.QtWidgets").QApplication
    app = QApplication.instance() or QApplication([])
    output = tmp_path / "record.pdf"
    ReportService().generate_pdf_record(
        {"id": 1, "name": "样品 A", "status": "FINISHED"}, [], [], output,
    )
    assert output.read_bytes().startswith(b"%PDF")
    assert output.stat().st_size > 1000
    app.quit()
