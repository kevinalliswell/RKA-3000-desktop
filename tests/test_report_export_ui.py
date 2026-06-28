from datetime import datetime, timedelta

import pytest

pytest.importorskip("PySide6")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from rka_desktop.core.config import Settings
from rka_desktop.core.modbus_client import ModbusRtuClient
from rka_desktop.db.database import DatabaseManager
from rka_desktop.models.entities import Measurement
from rka_desktop.ui.main_window import MainWindow


def _window(tmp_path):
    app = QApplication.instance() or QApplication([])
    settings = Settings(db_path=str(tmp_path / "report.db"), simulation_mode=True)
    db = DatabaseManager(settings.db_file)
    db.init_schema()
    exp_id = db.create_experiment("redox")
    db.set_experiment_sample_params(exp_id, 50.0, 60.0, 0.0)
    t0 = datetime(2026, 6, 28, 10, 0, 0)
    db.insert_measurements(
        [
            Measurement(
                experiment_id=exp_id,
                sampled_at=t0 + timedelta(minutes=i),
                temperature_c=900.0,
                weight_g=50.0 - 6.45 * i,
                pressure_mpa=0.0,
                n2_lpm=4.0,
                h2_lpm=1.0,
                co_lpm=0.0,
            )
            for i in range(3)
        ]
    )
    client = ModbusRtuClient(serial_port="loop://", baudrate=9600, unit_id=1, simulation_mode=True)
    return app, db, MainWindow(settings=settings, db=db, client=client), exp_id


def test_help_menu_and_report_button_state(tmp_path):
    app, _db, win, _exp = _window(tmp_path)
    menu_titles = [a.text() for a in win.menuBar().actions()]
    assert "帮助" in menu_titles
    assert win.analysis_page.btn_report.isEnabled() is False  # disabled until a compute
    win.close()
    app.quit()


def test_export_analysis_report_writes_file(tmp_path, monkeypatch):
    app, _db, win, exp_id = _window(tmp_path)
    page = win.analysis_page
    page.refresh_experiments()
    page.experiment_combo.setCurrentIndex(page.experiment_combo.findData(exp_id))
    page.spin_m0.setValue(50.0)
    page.spin_tfe.setValue(60.0)
    page.spin_feo.setValue(0.0)
    assert page.compute() is not None
    assert page.btn_report.isEnabled() is True

    out = tmp_path / "report.xlsx"
    monkeypatch.setattr(win.dialogs, "select_save_path", lambda *a, **k: out)
    # info/warning are modal QMessageBox dialogs that would block headlessly
    monkeypatch.setattr(win.dialogs, "info", lambda *a, **k: None)
    monkeypatch.setattr(win.dialogs, "warning", lambda *a, **k: None)
    win.export_analysis_report()

    assert out.exists()
    wb = load_workbook(out)
    assert "报告" in wb.sheetnames
    win.close()
    app.quit()
